#!/usr/bin/env python3
"""Impor sheet MATRIX (akademik) -> data/jadwal.json.  Pakai: python3 scripts/import_matrix.py file.xlsx [--write]"""
import openpyxl, re, json, sys, shutil, collections, datetime

PRODI = {"MG": "Metalurgi Pertambangan", "TMG": "Metalurgi Pertambangan", "TRM": "Teknik Rekayasa Multimedia",
         "TRP": "Teknologi Rekayasa Pangan", "ARS": "Arsitektur", "SPL": "Sipil", "ELK": "Teknik Elektronika", "MSN": "Mesin"}
HARI = {"SENIN": "Senin", "SELASA": "Selasa", "RABU": "Rabu", "KAMIS": "Kamis", "JUMAT": "Jumat", "SABTU": "Sabtu"}
clean = lambda v: re.sub(r"\s+", " ", str(v)).strip()

def jam_of(label):
    m = re.findall(r"(\d{1,2})[.:](\d{2})", str(label))
    return [f"{int(a):02d}:{b}" for a, b in m[:2]] if len(m) >= 2 else None

def main(path, write=False):
    wb = openpyxl.load_workbook(path, data_only=True); ws = wb["MATRIX"]
    dosen = {}
    for r in range(48, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, (int, float)) and b: dosen[int(a)] = clean(b)
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    blocks = []
    for hr in (2, 14, 26):
        for c in range(3, ws.max_column + 1):
            v = ws.cell(hr, c).value
            if v and clean(v).upper() in HARI:
                m = merged.get((hr, c)); blocks.append((hr, HARI[clean(v).upper()], c, m.max_col if m else c + 13))
    out, nodosen, unknown = [], [], set()
    for hr, hari, c1, c2 in blocks:
        for c in range(c1, c2 + 1):
            ruang = re.sub(r"\s*\(.*?\)", "", clean(ws.cell(hr + 1, c).value or "")).strip(" /")
            for r in range(hr + 2, hr + 11):
                v = ws.cell(r, c).value
                if v in (None, ""): continue
                s = clean(v)
                if len(s) < 3 or s.lower() == "no": continue
                m = merged.get((r, c)); r2 = m.max_row if m and m.max_col == c else r
                j1, j2 = jam_of(ws.cell(r, 2).value), jam_of(ws.cell(r2, 2).value)
                if not (j1 and j2): continue
                # kode dosen di akhir: (37) (35&19) (20/14)
                names, label = [], s
                mc = re.search(r"\((\d+(?:\s*[&/,]\s*\d+)*)\)?\s*$", s)
                if mc:
                    label = s[:mc.start()].strip()
                    for k in re.split(r"\s*[&/,]\s*", mc.group(1)):
                        if int(k) in dosen: names.append(dosen[int(k)])
                        else: unknown.add(int(k))
                else:
                    mn = re.search(r"\(([A-Za-z ]+)\)\s*$", s)
                    if mn and "GAB" not in mn.group(1).upper(): names, label = [mn.group(1).title()], s[:mn.start()].strip()
                if not names: nodosen.append((hari, ruang, s[:40]))
                # matkul / prodi / kelas
                prodi, matkul = "-", label
                if re.search(r"\bGAB\b", label, re.I): prodi = "Kelas Gabungan"
                else:
                    mk = re.search(r"\b(TMG|MG|TRM|TRP|ARS|SPL|ELK|MSN)\b\s*([A-Z.]*\s*[\d.]+[A-Z.]*.*)$", label)
                    if mk:
                        prodi = f"{PRODI[mk.group(1)]} {mk.group(2).strip()}"
                        matkul = label[:mk.start()].strip() or label
                out.append({"hari": hari, "jam": f"{j1[0]} - {j2[1]}", "matkul": matkul, "prodi": prodi,
                            "dosen": " / ".join(names) if names else "-", "ruangan": ruang, "status": "Belum Mulai"})
    dup = collections.Counter((x["hari"], x["jam"], x["ruangan"]) for x in out)
    print("total:", len(out), "| per hari:", dict(collections.Counter(x["hari"] for x in out)))
    print("tanpa dosen:", len(nodosen), "| kode dosen tak dikenal:", sorted(unknown))
    print("ruangan:", sorted({x["ruangan"] for x in out}))
    print("bentrok ruang+jam:", [k for k, v in dup.items() if v > 1][:8])
    if write:
        shutil.copy("data/jadwal.json", "data/jadwal.backup-%s.json" % datetime.date.today())
        json.dump(out, open("data/jadwal.json", "w", encoding="utf8"), ensure_ascii=False, indent=2)
        print("ditulis ke data/jadwal.json (cadangan lama tersimpan)")
    return out, nodosen

if __name__ == "__main__":
    o, nd = main(sys.argv[1], "--write" in sys.argv)
    for x in o[:4] + o[300:303]: print(x)
    print(nd[:10])
